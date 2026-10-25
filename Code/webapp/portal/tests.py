"""Web app tests: sign-in, roles, uploads, search output, document actions, downloads and CSRF.

Run with ``cd Code/webapp && python manage.py test``. Everything runs in local mode against a temporary data
folder, with documents processed inline (no background thread).
"""

import io
import os
import shutil
import tempfile
from pathlib import Path
from unittest import mock

from django.core.cache import cache
from django.test import Client, SimpleTestCase, override_settings
from django.urls import reverse

from advancesearch.config import reset_settings_cache
from advancesearch.ingest import doc_id_for

PASSWORD = "correct horse battery staple"


def make_pdf(pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        c.drawString(50, 800, text)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(paragraphs):
    from docx import Document

    d = Document()
    d.add_heading("Guide", level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def upload_file(name, data):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(name, data)


@override_settings(AUTH_MODE="password", DEMO_PASSWORD=PASSWORD)
class AppTestCase(SimpleTestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="advsearch-web-")
        env = {"ADVSEARCH_MODE": "local", "LOCAL_DATA_DIR": self.tmp, "INGEST_WORKER": "inline",
               "EMBEDDING_DIMENSIONS": "256", "MAX_UPLOAD_MB": "20"}
        self.env = mock.patch.dict(os.environ, env)
        self.env.start()
        for name in ("STORAGE_BACKEND", "SEARCH_BACKEND", "EMBEDDING_BACKEND"):
            os.environ.pop(name, None)
        reset_settings_cache()
        from portal import data

        data.reset_index_status()
        cache.clear()

    def tearDown(self):
        reset_settings_cache()
        self.env.stop()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def sign_in(self, role="Admin", client=None):
        client = client or self.client
        response = client.post(reverse("accounts:login"), {"role": role, "password": PASSWORD})
        self.assertEqual(response.status_code, 302)
        return client

    def upload(self, name, data, client=None):
        client = client or self.client
        return client.post(reverse("portal:upload"), {"file": upload_file(name, data)}, HTTP_ACCEPT="application/json")

    def seed(self):
        self.sign_in("Admin")
        docs = {
            "Leave-Policy.pdf": make_pdf(["Every employee receives an annual entitlement of 24 working days of leave.",
                                          "Sick leave is 12 days per year with a medical certificate."]),
            "Tablet-Spec.pdf": make_pdf(["The tablet battery lasts 14 hours on one charge."]),
            "Security.docx": make_docx(["Report a lost or stolen laptop to the service desk within one hour."]),
            "Retention.pdf": make_pdf(["Invoices are kept for eight years, CCTV footage for 30 days."]),
        }
        for name, data in docs.items():
            self.assertEqual(self.upload(name, data).status_code, 200)
        return docs


# ---------------------------------------------------------------------------- sign-in

class DemoPasswordTests(AppTestCase):
    def test_login_page_shows_password_and_roles(self):
        response = self.client.get(reverse("accounts:login"))
        self.assertContains(response, 'name="password"')
        self.assertContains(response, 'value="Admin"')
        self.assertContains(response, 'value="User"')

    def test_wrong_password_403_then_429_after_five_attempts(self):
        for attempt in range(5):
            response = self.client.post(reverse("accounts:login"), {"role": "Admin", "password": "nope"})
            self.assertEqual(response.status_code, 403, attempt)
            self.assertContains(response, "Wrong password", status_code=403)
        response = self.client.post(reverse("accounts:login"), {"role": "Admin", "password": PASSWORD})
        self.assertEqual(response.status_code, 429)  # blocked even with the right password
        self.assertContains(response, "15 minutes", status_code=429)
        self.assertNotIn("user", self.client.session)

    def test_lockout_is_per_client_ip_from_forwarded_for(self):
        for _ in range(5):
            self.client.post(reverse("accounts:login"), {"role": "User", "password": "x"},
                             HTTP_X_FORWARDED_FOR="203.0.113.7:51234")
        blocked = self.client.post(reverse("accounts:login"), {"role": "User", "password": PASSWORD},
                                   HTTP_X_FORWARDED_FOR="203.0.113.7:51234")
        self.assertEqual(blocked.status_code, 429)
        other = Client().post(reverse("accounts:login"), {"role": "User", "password": PASSWORD},
                              HTTP_X_FORWARDED_FOR="198.51.100.9:4000")
        self.assertEqual(other.status_code, 302)

    def test_forged_forwarded_for_prefix_does_not_reset_the_counter(self):
        for i in range(5):
            self.client.post(reverse("accounts:login"), {"role": "User", "password": "x"},
                             HTTP_X_FORWARDED_FOR=f"10.0.0.{i}, 203.0.113.7:5000")
        response = self.client.post(reverse("accounts:login"), {"role": "User", "password": PASSWORD},
                                    HTTP_X_FORWARDED_FOR="10.9.9.9, 203.0.113.7:6000")
        self.assertEqual(response.status_code, 429)

    def test_right_password_signs_in_with_chosen_role(self):
        for role in ("Admin", "User"):
            client = Client()
            response = client.post(reverse("accounts:login"), {"role": role, "password": PASSWORD})
            self.assertRedirects(response, reverse("portal:search"), fetch_redirect_response=False)
            self.assertEqual(client.session["user"]["role"], role)
            page = client.get(reverse("portal:search"))
            self.assertContains(page, f'role-{role.lower()}')

    def test_unknown_role_is_rejected(self):
        response = self.client.post(reverse("accounts:login"), {"role": "Owner", "password": PASSWORD})
        self.assertEqual(response.status_code, 400)

    @override_settings(DEMO_PASSWORD="")
    def test_empty_configured_password_refuses_sign_in(self):
        page = self.client.get(reverse("accounts:login"))
        self.assertContains(page, "not configured")
        self.assertNotContains(page, 'name="password"')
        for given in ("", "anything"):
            response = self.client.post(reverse("accounts:login"), {"role": "Admin", "password": given})
            self.assertEqual(response.status_code, 503)
        self.assertNotIn("user", self.client.session)

    @override_settings(AUTH_MODE="dev")
    def test_dev_mode_needs_no_password(self):
        response = self.client.post(reverse("accounts:login"), {"role": "User"})
        self.assertEqual(response.status_code, 302)

    def test_next_parameter_cannot_redirect_off_site(self):
        response = self.client.post(reverse("accounts:login") + "?next=https://evil.example/",
                                    {"role": "User", "password": PASSWORD})
        self.assertEqual(response["Location"], reverse("portal:search"))
        response = Client().post(reverse("accounts:login") + "?next=/settings", {"role": "User", "password": PASSWORD})
        self.assertEqual(response["Location"], "/settings")

    def test_logout_requires_post_and_clears_session(self):
        self.sign_in("User")
        self.assertEqual(self.client.get(reverse("accounts:logout")).status_code, 405)
        self.client.post(reverse("accounts:logout"))
        self.assertEqual(self.client.get(reverse("portal:search")).status_code, 302)

    def test_tampered_session_role_is_ignored(self):
        session = self.client.session
        session["user"] = {"role": "Superuser"}
        session.save()
        self.assertEqual(self.client.get(reverse("portal:search")).status_code, 302)


class UnauthenticatedTests(AppTestCase):
    def test_pages_redirect_to_sign_in(self):
        for url in ("/", "/documents", "/settings", "/search/results?q=x", f"/documents/{'a' * 32}/open/x.pdf"):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 302, url)
            self.assertTrue(response["Location"].startswith("/auth/login?next="), url)

    def test_htmx_requests_get_hx_redirect(self):
        response = self.client.get("/search/results?q=x", HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 401)
        self.assertTrue(response["HX-Redirect"].startswith("/auth/login"))

    def test_api_returns_401(self):
        for url in ("/api/search?q=x", "/api/documents", "/api/status"):
            self.assertEqual(self.client.get(url).status_code, 401, url)

    def test_post_actions_redirect(self):
        for url in ("/documents/upload", "/documents/sync", "/documents/rebuild", f"/documents/{'a' * 32}/delete"):
            self.assertEqual(self.client.post(url).status_code, 302, url)

    def test_healthz_is_public(self):
        self.assertEqual(self.client.get("/healthz").content, b"ok")


# ---------------------------------------------------------------------------- roles

class RoleTests(AppTestCase):
    ADMIN_GETS = ["/documents", "/documents/table", "/settings"]
    ADMIN_POSTS = ["/documents/upload", "/documents/sync", "/documents/rebuild", f"/documents/{'a' * 32}/reindex",
                   f"/documents/{'a' * 32}/delete", "/settings"]
    ADMIN_APIS = ["/api/documents", "/api/status"]

    def test_user_cannot_reach_admin_pages_or_apis(self):
        self.sign_in("User")
        for url in self.ADMIN_GETS:
            self.assertEqual(self.client.get(url).status_code, 403, url)
        for url in self.ADMIN_POSTS:
            self.assertEqual(self.client.post(url).status_code, 403, url)
        for url in self.ADMIN_APIS:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 403, url)
            self.assertEqual(response["Content-Type"], "application/json")

    def test_user_navigation_hides_admin_pages(self):
        self.sign_in("User")
        page = self.client.get("/")
        self.assertNotContains(page, 'href="/documents"')
        self.assertNotContains(page, 'href="/settings"')

    def test_user_can_search_and_open(self):
        docs = self.seed()
        user = self.sign_in("User", Client())
        self.assertEqual(user.get("/").status_code, 200)
        self.assertEqual(user.get("/api/search?q=annual leave").status_code, 200)
        name = "Leave-Policy.pdf"
        response = user.get(reverse("portal:open", args=[doc_id_for(name), name]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(b"".join(response.streaming_content), docs[name])

    def test_admin_can_reach_everything(self):
        self.sign_in("Admin")
        for url in self.ADMIN_GETS + self.ADMIN_APIS + ["/"]:
            self.assertEqual(self.client.get(url).status_code, 200, url)
        page = self.client.get("/")
        self.assertContains(page, 'href="/documents"')
        self.assertContains(page, 'href="/settings"')


# ---------------------------------------------------------------------------- uploads

class UploadTests(AppTestCase):
    def setUp(self):
        super().setUp()
        self.sign_in("Admin")

    def test_valid_upload_is_indexed(self):
        response = self.upload("Leave Policy.pdf", make_pdf(["Annual leave entitlement is 24 days."]))
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertTrue(body["ok"])
        self.assertEqual(body["name"], "Leave Policy.pdf")
        docs = self.client.get("/api/documents").json()["documents"]
        self.assertEqual([(d["fileName"], d["status"], d["uploadedBy"]) for d in docs],
                         [("Leave Policy.pdf", "indexed", "Admin")])

    def test_wrong_type_is_rejected(self):
        response = self.upload("notes.txt", b"plain text")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Only PDF", response.json()["message"])
        self.assertEqual(self.upload("old.doc", b"\xd0\xcf\x11\xe0").status_code, 400)

    def test_wrong_magic_bytes_are_rejected(self):
        response = self.upload("fake.pdf", b"MZ\x90\x00 an executable")
        self.assertEqual(response.status_code, 400)
        self.assertIn("not a valid .pdf", response.json()["message"])
        response = self.upload("fake.xlsx", make_docx(["a word file renamed"]))
        self.assertIn("not a valid .xlsx", response.json()["message"])

    def test_too_large_is_rejected(self):
        with mock.patch.dict(os.environ, {"MAX_UPLOAD_MB": "1"}):
            reset_settings_cache()
            response = self.upload("big.pdf", b"%PDF-1.4\n" + b"0" * (1024 * 1024 + 10))
        self.assertEqual(response.status_code, 400)
        self.assertIn("1 MB limit", response.json()["message"])

    def test_odd_file_names_are_sanitised(self):
        for given, stored in [("../../secret.pdf", "secret.pdf"), ("Résumé <final>.docx", "Resume_final.docx"),
                              ("a\\b\\c.xlsx", "c.xlsx")]:
            data = make_docx(["x"]) if stored.endswith(".docx") else make_pdf([given]) if stored.endswith(".pdf") else None
            if stored.endswith(".xlsx"):
                from openpyxl import Workbook

                wb, buf = Workbook(), io.BytesIO()
                wb.active.append(["Header"])
                wb.active.append([given])
                wb.save(buf)
                data = buf.getvalue()
            response = self.upload(given, data)
            self.assertEqual(response.json().get("name"), stored, response.json())
        names = sorted(d["fileName"] for d in self.client.get("/api/documents").json()["documents"])
        self.assertEqual(names, ["Resume_final.docx", "c.xlsx", "secret.pdf"])

    def test_duplicate_content_says_already_indexed(self):
        data = make_pdf(["Same content."])
        self.upload("one.pdf", data)
        response = self.upload("two.pdf", data)
        self.assertEqual(response.status_code, 400)
        self.assertIn("already indexed as one.pdf", response.json()["message"])

    def test_scanned_pdf_shows_no_text(self):
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.rect(10, 10, 100, 100, fill=1)
        c.showPage()
        c.save()
        self.upload("scan.pdf", buf.getvalue())
        table = self.client.get("/documents/table")
        self.assertContains(table, "No text")
        self.assertContains(table, "No searchable text found. This file may be a scan.")

    def test_upload_without_javascript_redirects_with_messages(self):
        response = self.client.post("/documents/upload", {"file": [upload_file("a.pdf", make_pdf(["A"])),
                                                                   upload_file("b.txt", b"no")]})
        self.assertRedirects(response, "/documents", fetch_redirect_response=False)
        page = self.client.get("/documents")
        self.assertContains(page, "a.pdf: Queued for indexing.")
        self.assertContains(page, "b.txt: Only PDF")

    def test_no_file(self):
        response = self.client.post("/documents/upload", HTTP_ACCEPT="application/json")
        self.assertEqual(response.status_code, 400)


# ---------------------------------------------------------------------------- search

class SearchTests(AppTestCase):
    def test_no_documents_state(self):
        self.sign_in("Admin")
        response = self.client.get("/search/results?q=anything", HTTP_HX_REQUEST="true")
        self.assertContains(response, "No documents indexed yet")
        self.assertContains(response, "Upload documents")
        user = self.sign_in("User", Client())
        response = user.get("/search/results?q=anything", HTTP_HX_REQUEST="true")
        self.assertContains(response, "An Admin needs to upload")
        self.assertNotContains(response, "Upload documents")

    def test_api_search_returns_top_3_with_scores_and_highlights(self):
        self.seed()
        body = self.client.get("/api/search", {"q": "how many days off do I get"}).json()
        self.assertLessEqual(body["count"], 3)
        self.assertEqual(body["topN"], 3)
        top = body["results"][0]
        self.assertEqual(top["fileName"], "Leave-Policy.pdf")
        self.assertEqual(top["rank"], 1)
        self.assertEqual(top["scoreSource"], "semantic")
        self.assertTrue(0 <= top["matchPercent"] <= 100)
        self.assertEqual(top["matchPercent"], round(top["rerankerScore"] / 4 * 100))
        self.assertIn("leave", [h.lower() for h in top["highlights"]])
        self.assertTrue(top["openUrl"].startswith("http://testserver/documents/"))
        self.assertEqual(top["location"], "Page 1")
        scores = [r["matchPercent"] for r in body["results"]]
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertTrue(all(s >= body["minMatchPercent"] for s in scores))

    def test_results_partial_renders_cards(self):
        self.seed()
        response = self.client.get("/search/results", {"q": "battery life of the tablet"}, HTTP_HX_REQUEST="true")
        self.assertContains(response, 'class="result-card"')
        self.assertContains(response, "Tablet-Spec.pdf")
        self.assertContains(response, "<mark>")
        self.assertContains(response, 'role="meter"')
        self.assertContains(response, "Open document")
        self.assertContains(response, "Copy link")
        self.assertContains(response, "How is the match % calculated?")
        self.assertContains(response, "not a probability")

    def test_full_page_with_query_renders_results(self):
        self.seed()
        response = self.client.get("/", {"q": "lost laptop"})
        self.assertContains(response, "Security.docx")
        self.assertContains(response, "has-query")

    def test_non_htmx_results_request_redirects_to_page(self):
        self.sign_in("User")
        response = self.client.get("/search/results", {"q": "a b"})
        self.assertRedirects(response, "/?q=a%20b", fetch_redirect_response=False)

    def test_no_close_match_state(self):
        self.seed()
        response = self.client.get("/search/results", {"q": "football world cup"}, HTTP_HX_REQUEST="true")
        self.assertContains(response, "No close match")

    def test_empty_query(self):
        self.seed()
        response = self.client.get("/search/results", {"q": "   "}, HTTP_HX_REQUEST="true")
        self.assertNotContains(response, "result-card")
        self.assertEqual(self.client.get("/api/search?q=").status_code, 400)

    def test_query_and_snippets_are_escaped(self):
        self.seed()
        self.upload("x.pdf", make_pdf(["<script>alert(1)</script> annual leave rules"]))
        response = self.client.get("/search/results", {"q": "<img src=x onerror=alert(1)> annual leave"},
                                   HTTP_HX_REQUEST="true")
        self.assertNotContains(response, "<script>alert(1)</script>")
        self.assertNotContains(response, "<img src=x")
        self.assertContains(response, "&lt;script&gt;")

    def test_settings_change_top_n_and_minimum(self):
        self.seed()
        self.client.post("/settings", {"top_n": "1", "min_match_percent": "0", "example_queries": "", "welcome_text": ""})
        self.assertEqual(self.client.get("/api/search?q=leave").json()["count"], 1)
        self.client.post("/settings", {"top_n": "3", "min_match_percent": "100", "example_queries": "", "welcome_text": ""})
        self.assertEqual(self.client.get("/api/search?q=leave").json()["count"], 0)

    def test_search_failure_shows_error_state(self):
        self.seed()
        with mock.patch("portal.data.run_search", side_effect=RuntimeError("index unreachable")):
            response = self.client.get("/search/results", {"q": "leave"}, HTTP_HX_REQUEST="true")
            self.assertContains(response, "Search is temporarily unavailable")
            self.assertContains(response, "index unreachable")  # admins see the detail
            self.assertEqual(self.client.get("/api/search?q=leave").status_code, 503)
            user = self.sign_in("User", Client())
            self.assertNotContains(user.get("/search/results", {"q": "leave"}, HTTP_HX_REQUEST="true"),
                                   "index unreachable")


# ---------------------------------------------------------------------------- document actions

class DocumentActionTests(AppTestCase):
    def test_delete_removes_document_and_chunks(self):
        self.seed()
        from advancesearch import factory

        name = "Leave-Policy.pdf"
        before = factory.get_index().document_chunk_count(doc_id_for(name))
        self.assertEqual(before, 2)
        response = self.client.post(reverse("portal:delete", args=[doc_id_for(name)]), {"name": name},
                                    HTTP_HX_REQUEST="true")
        self.assertContains(response, "Deleted Leave-Policy.pdf and its 2 passages.")
        self.assertEqual(factory.get_index().document_chunk_count(doc_id_for(name)), 0)
        self.assertNotIn(name, [d["fileName"] for d in self.client.get("/api/documents").json()["documents"]])
        results = self.client.get("/api/search?q=annual leave entitlement").json()["results"]
        self.assertNotIn(name, [r["fileName"] for r in results])

    def test_reindex(self):
        self.seed()
        name = "Security.docx"
        response = self.client.post(reverse("portal:reindex", args=[doc_id_for(name)]), {"name": name})
        self.assertRedirects(response, "/documents", fetch_redirect_response=False)
        doc = next(d for d in self.client.get("/api/documents").json()["documents"] if d["fileName"] == name)
        self.assertEqual(doc["status"], "indexed")

    def test_action_with_mismatched_name_is_404(self):
        self.seed()
        response = self.client.post(reverse("portal:delete", args=[doc_id_for("Security.docx")]), {"name": "Leave-Policy.pdf"})
        self.assertEqual(response.status_code, 404)
        self.assertEqual(len(self.client.get("/api/documents").json()["documents"]), 4)

    def test_sync_and_rebuild(self):
        self.seed()
        response = self.client.post("/documents/sync", HTTP_HX_REQUEST="true")
        self.assertContains(response, "every document is already indexed")
        response = self.client.post("/documents/rebuild", HTTP_HX_REQUEST="true")
        self.assertContains(response, "4 documents are queued")
        statuses = {d["status"] for d in self.client.get("/api/documents").json()["documents"]}
        self.assertEqual(statuses, {"indexed"})

    def test_documents_page_filters(self):
        self.seed()
        table = self.client.get("/documents/table", {"type": "docx"})
        self.assertContains(table, "Security.docx")
        self.assertNotContains(table, "Leave-Policy.pdf")
        table = self.client.get("/documents/table", {"q": "leave"})
        self.assertContains(table, "Leave-Policy.pdf")
        self.assertNotContains(table, "Security.docx")
        table = self.client.get("/documents/table", {"status": "failed"})
        self.assertContains(table, "No documents match")


class DownloadTests(AppTestCase):
    def test_open_streams_original_file(self):
        docs = self.seed()
        for name, ctype, disposition in [("Leave-Policy.pdf", "application/pdf", "inline"),
                                         ("Security.docx", "application/vnd.openxmlformats-officedocument."
                                                           "wordprocessingml.document", "attachment")]:
            response = self.client.get(reverse("portal:open", args=[doc_id_for(name), name]))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], ctype)
            self.assertTrue(response["Content-Disposition"].startswith(f'{disposition}; filename="{name}"'))
            self.assertEqual(response["Content-Length"], str(len(docs[name])))
            self.assertEqual(response["Cache-Control"], "private, no-store")
            self.assertEqual(b"".join(response.streaming_content), docs[name])

    def test_open_with_wrong_id_or_missing_file_is_404(self):
        self.seed()
        self.assertEqual(self.client.get(reverse("portal:open", args=["0" * 32, "Leave-Policy.pdf"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("portal:open", args=[doc_id_for("nope.pdf"), "nope.pdf"])).status_code, 404)
        self.assertEqual(self.client.get(reverse("portal:open", args=[doc_id_for("../x.pdf"), "../x.pdf"])).status_code, 404)


# ---------------------------------------------------------------------------- settings, CSRF, headers

class SettingsPageTests(AppTestCase):
    def test_save_and_reset(self):
        self.sign_in("Admin")
        response = self.client.post("/settings", {"top_n": "5", "min_match_percent": "30",
                                                  "example_queries": "first\nsecond", "welcome_text": "Hello there"})
        self.assertRedirects(response, "/settings", fetch_redirect_response=False)
        page = self.client.get("/")
        self.assertContains(page, "Hello there")
        self.assertContains(page, 'data-query="second"')
        self.client.post("/settings", {"action": "reset"})
        self.assertNotContains(self.client.get("/"), "Hello there")

    def test_invalid_values_show_error(self):
        self.sign_in("Admin")
        response = self.client.post("/settings", {"top_n": "11", "min_match_percent": "20", "example_queries": "",
                                                  "welcome_text": ""})
        self.assertContains(response, "between 1 and 10", status_code=400)


class CsrfAndHeaderTests(AppTestCase):
    def test_login_without_csrf_token_is_rejected(self):
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.post(reverse("accounts:login"), {"role": "Admin", "password": PASSWORD}).status_code, 403)

    def test_actions_without_csrf_token_are_rejected(self):
        self.seed()
        client = Client(enforce_csrf_checks=True)
        client.cookies = self.client.cookies
        name = "Leave-Policy.pdf"
        self.assertEqual(client.post(reverse("portal:delete", args=[doc_id_for(name)]), {"name": name}).status_code, 403)
        self.assertEqual(client.post("/documents/upload", {"file": upload_file("z.pdf", make_pdf(["z"]))}).status_code, 403)
        self.assertIn(name, [d["fileName"] for d in self.client.get("/api/documents").json()["documents"]])

    def test_csrf_token_works_for_real_forms(self):
        client = Client(enforce_csrf_checks=True)
        page = client.get(reverse("accounts:login"))
        token = page.context["csrf_token"]
        response = client.post(reverse("accounts:login"), {"role": "User", "password": PASSWORD,
                                                           "csrfmiddlewaretoken": str(token)})
        self.assertEqual(response.status_code, 302)

    def test_security_headers(self):
        response = self.client.get(reverse("accounts:login"))
        csp = response["Content-Security-Policy"]
        self.assertIn("script-src 'self'", csp)
        self.assertNotIn("unsafe-inline", csp)
        self.assertIn("frame-ancestors 'none'", csp)
        self.assertEqual(response["X-Frame-Options"], "DENY")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")

    def test_templates_have_no_inline_scripts_or_styles(self):
        import re

        root = Path(__file__).resolve().parent.parent / "templates"
        for path in root.rglob("*.html"):
            text = path.read_text()
            self.assertNotRegex(text, r"<script(?![^>]*\bsrc=)", path.name)
            self.assertNotRegex(text, r"\sstyle=\"", path.name)
            self.assertIsNone(re.search(r"\son[a-z]+=\"", text), path.name)

    def test_error_pages(self):
        self.sign_in("User")
        self.assertContains(self.client.get("/nope"), "Page not found", status_code=404)
        self.assertContains(self.client.get("/settings"), "Not available for your role", status_code=403)
        self.assertEqual(self.client.get("/api/nope").json(), {"error": "Not found."})


class StaticFilesTests(SimpleTestCase):
    def test_front_end_libraries_are_vendored(self):
        root = Path(__file__).resolve().parent.parent
        self.assertTrue((root / "static" / "vendor" / "htmx-2.0.4.min.js").is_file())
        base = (root / "templates" / "base.html").read_text()
        self.assertNotIn("https://", base.replace("https://", "", 0) if False else base.split("<body")[0])
