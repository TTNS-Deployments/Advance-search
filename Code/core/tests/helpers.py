"""Small document builders for tests (PDF, DOCX, XLSX written to a temp folder)."""
from __future__ import annotations

import io
from pathlib import Path


def make_pdf(path: Path, pages: list[str]) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            if line:
                c.drawString(50, y, line)
            y -= 16
        c.showPage()
    c.save()
    return path


def make_blank_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    c.rect(50, 50, 200, 200, fill=1)
    c.showPage()
    c.save()
    return path


def make_docx(path: Path, blocks: list[tuple]) -> Path:
    """blocks: ("h1", text) / ("h2", text) / ("p", text) / ("table", rows)"""
    from docx import Document

    d = Document()
    for kind, value in blocks:
        if kind.startswith("h"):
            d.add_heading(value, level=int(kind[1:]))
        elif kind == "p":
            d.add_paragraph(value)
        elif kind == "table":
            table = d.add_table(rows=0, cols=len(value[0]))
            for row in value:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
    d.save(str(path))
    return path


def make_xlsx(path: Path, sheets: dict[str, list[list]], hidden: tuple[str, ...] = ()) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
        if title in hidden:
            ws.sheet_state = "hidden"
    wb.save(str(path))
    return path


def upload(service, path: Path, by: str = "Admin", process: bool = True):
    """Upload like the Documents page does; ``process`` also runs the ingestion step at once."""
    with path.open("rb") as fh:
        result = service.upload(path.name, fh, path.stat().st_size, uploaded_by=by)
    if process:
        service.process(result.name)
    return result


def bytes_io(data: bytes) -> io.BytesIO:
    return io.BytesIO(data)


def png(width: int, height: int, colour: str = "#2f4a8a", label: str = "") -> bytes:
    """A PNG with a shape on it (each colour/label gives different bytes)."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((2, 2, width - 3, height - 3), outline=colour, width=2)
    if label:
        draw.text((6, height // 2 - 6), label, fill="black")
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def make_pdf_with_images(path: Path, pages: list[tuple[list[str], list[tuple[bytes, int, int, int, int]]]]) -> Path:
    """pages: (lines of text from the top, images as (png, x, y, width, height) in points)."""
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for lines, images in pages:
        y = 800
        for line in lines:
            if line:
                c.drawString(50, y, line)
            y -= 20
        for data, x, iy, w, h in images:
            c.drawImage(ImageReader(io.BytesIO(data)), x, iy, w, h)
        c.showPage()
    c.save()
    return path


def make_docx_with_image(path: Path, before: str, image: bytes, name: str, alt: str, caption: str, after: str) -> Path:
    from docx import Document
    from docx.shared import Cm

    d = Document()
    d.add_heading("Network design", level=1)
    d.add_paragraph(before)
    d.add_picture(io.BytesIO(image), width=Cm(8))
    inline = d.inline_shapes[-1]._inline
    inline.docPr.set("descr", alt)
    inline.graphic.graphicData.pic.nvPicPr.cNvPr.set("name", name)
    d.add_paragraph(caption, style="Caption")
    d.add_paragraph(after)
    d.save(str(path))
    return path


def make_xlsx_with_image(path: Path, rows: list[list], image: bytes, anchor: str = "A4") -> Path:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image

    wb = Workbook()
    ws = wb.active
    ws.title = "Plan"
    for row in rows:
        ws.append(row)
    ws.add_image(Image(io.BytesIO(image)), anchor)
    wb.save(str(path))
    return path
