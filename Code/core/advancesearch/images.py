"""Embedded images: extraction from PDF (pypdf), Word and Excel files (their zip parts), filtering, de-duplication,
the searchable text that describes each image, and the image store.

Each image that is kept gets a stable id (the first 16 hex characters of the SHA-256 of its bytes), is saved as
``<doc_id>/<image_id>.<ext>`` in the images container (``.localdata/storage/extracted-images`` locally) and is
listed with its metadata in ``<doc_id>/images.json``. Its text, built from the original file name, alt text, the
figure caption, the text just above and below it and (optionally) OCR, is indexed as one passage of the document,
so a query such as "architecture diagram", "figure 3" or an image's file name finds the image.

Skipped: images under ``min_width`` x ``min_height`` pixels (icons, bullets), thin rules and borders (aspect ratio
above ``max_aspect_ratio``), the same image repeated on ``repeated_pages`` or more pages (logos, watermarks),
formats a browser cannot show and cannot be converted (WMF, EMF), and copies of an image already kept.
"""

from __future__ import annotations

import hashlib
import io
import logging
import posixpath
import re
import threading
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator
from xml.etree import ElementTree as ET

log = logging.getLogger(__name__)

WEB_FORMATS = {"png": "image/png", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp"}
MANIFEST = "images.json"
IMAGE_ID_RE = re.compile(r"[0-9a-f]{16}")

# "Figure 3", "Fig. 2a", "Table IV", "Diagram 1:" at the start of a line
CAPTION_RE = re.compile(r"^\s*(?:fig(?:ure)?|table|chart|diagram|image|photo|picture|exhibit|illustration|graph|"
                        r"map|plate|screenshot|drawing)\.?\s*(?:\d+[a-z]?|[ivxlc]+|[A-Z])\b", re.IGNORECASE)
# Names that say nothing about the image: "Im1.png", "image3.jpeg", "Picture 2", "FormXob.1f4d...png"
GENERIC_NAME_RE = re.compile(r"^(?:im|img|image|images|picture|pic|photo|x|obj|object|graphic|formxob|media|"
                             r"grafik|bild)?[\s._-]*(?:[0-9a-f]{8,}|\d+)?$", re.IGNORECASE)

NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "v": "urn:schemas-microsoft-com:vml",
    "o": "urn:schemas-microsoft-com:office:office",
}


def _q(prefix: str, tag: str) -> str:
    return f"{{{NS[prefix]}}}{tag}"


@dataclass(frozen=True)
class ImageOptions:
    enabled: bool = True
    min_width: int = 50
    min_height: int = 50
    max_aspect_ratio: float = 12.0
    repeated_pages: int = 3            # 0 = never treat a repeated image as decorative
    max_per_document: int = 100
    context_lines: int = 3             # lines (or paragraphs, or rows) taken above and below the image
    max_text_chars: int = 2000
    ocr: str = "auto"                  # auto | off

    @classmethod
    def from_config(cls, cfg: dict, enabled: bool = True, ocr: str = "auto") -> "ImageOptions":
        return cls(
            enabled=enabled and bool(cfg.get("enabled", True)), min_width=int(cfg.get("min_width", 50)),
            min_height=int(cfg.get("min_height", 50)), max_aspect_ratio=float(cfg.get("max_aspect_ratio", 12.0)),
            repeated_pages=int(cfg.get("repeated_pages", 3)), max_per_document=int(cfg.get("max_per_document", 100)),
            context_lines=int(cfg.get("context_lines", 3)), max_text_chars=int(cfg.get("max_text_chars", 2000)),
            ocr="off" if ocr == "off" else str(cfg.get("ocr", "auto")).lower(),
        )

    def keeps(self, width: int, height: int) -> bool:
        if width < self.min_width or height < self.min_height:
            return False
        return max(width, height) / max(1, min(width, height)) <= self.max_aspect_ratio


@dataclass
class _Candidate:
    """An image found in a file, before it is decoded, filtered and de-duplicated."""
    data: bytes
    location: str
    page: int | None = None
    original_name: str = ""
    alt_text: str = ""
    caption: str = ""
    nearby: list[str] = field(default_factory=list)
    shown_size: tuple[float, float] | None = None   # size on the PDF page in points, when known


@dataclass
class ExtractedImage:
    data: bytes
    format: str                        # png, jpeg, gif or webp
    width: int
    height: int
    location: str                      # "Page 3", "Section: Wiring", "Sheet 'Plan', cell B4"
    page: int | None = None
    original_name: str = ""
    alt_text: str = ""
    caption: str = ""
    nearby: str = ""
    ocr_text: str = ""
    occurrences: int = 1

    @property
    def image_id(self) -> str:
        return hashlib.sha256(self.data).hexdigest()[:16]

    @property
    def extension(self) -> str:
        return "jpg" if self.format == "jpeg" else self.format

    @property
    def content_type(self) -> str:
        return WEB_FORMATS[self.format]

    def search_text(self, max_chars: int = 2000, source: str = "") -> str:
        """The text indexed for this image; empty when nothing describes it (it is then stored but not indexed).
        ``source`` (the document's file name) gives the image the document's subject ("FT-200 field tablet")."""
        name = self.original_name if meaningful_name(self.original_name) else ""
        described = [t for t in (self.caption, self.alt_text, self.ocr_text, self.nearby) if t]
        if not name and not described:
            return ""
        where = f"{self.location} of {source}" if source else self.location
        lines = [f"Image: {name} ({where})" if name else f"Image ({where})"]
        seen = {""}
        for text in (self.caption, self.alt_text, self.ocr_text):
            if text.lower() not in seen:
                lines.append(text)
                seen.add(text.lower())
        described_text = " ".join(seen)
        nearby = [line for line in self.nearby.split("\n")
                  if line.strip() and line.strip().lower() not in described_text]  # not the caption again
        lines.extend(nearby)
        return "\n".join(lines)[:max_chars]

    def record(self, doc_id: str, file_name: str, path: str) -> dict:
        return {
            "imageId": self.image_id, "docId": doc_id, "fileName": file_name, "page": self.page,
            "location": self.location, "path": path, "originalName": self.original_name, "width": self.width,
            "height": self.height, "format": self.format, "bytes": len(self.data), "caption": self.caption,
            "altText": self.alt_text, "ocr": bool(self.ocr_text), "occurrences": self.occurrences,
        }


def meaningful_name(name: str) -> bool:
    stem = posixpath.basename(name or "").rsplit(".", 1)[0]
    return bool(stem.strip()) and not GENERIC_NAME_RE.fullmatch(stem.strip())


def _clean(text: str) -> str:
    return " ".join(str(text or "").split())


def _caption_of(lines: list[str]) -> str:
    return next((line for line in lines if CAPTION_RE.match(line)), "")


# ---------------------------------------------------------------------------- decoding

def pillow_available() -> bool:
    try:
        import PIL  # noqa: F401
    except ImportError:
        return False
    return True


def _decode(data: bytes) -> tuple[bytes, str, int, int] | None:
    """``(bytes, format, width, height)`` in a format browsers show, converting others to PNG; None if unreadable."""
    from PIL import Image

    try:
        with Image.open(io.BytesIO(data)) as im:
            fmt = (im.format or "").lower()
            width, height = im.size
            if fmt in WEB_FORMATS and im.mode != "CMYK":
                return data, fmt, width, height
            im.load()
            converted = im.convert("RGBA" if im.mode in ("RGBA", "LA", "P", "PA") else "RGB")
            out = io.BytesIO()
            converted.save(out, "PNG", optimize=True)
            return out.getvalue(), "png", width, height
    except Exception:  # noqa: BLE001 - unreadable, vector-only (WMF/EMF) or a decompression bomb: skip it
        return None


# ---------------------------------------------------------------------------- OCR (optional)

_ocr_lock = threading.Lock()
_ocr_available: bool | None = None


def ocr_text(data: bytes, mode: str = "auto") -> str:
    """Text in the image via Tesseract (``pytesseract``), or "" when OCR is off or not installed."""
    global _ocr_available
    if mode == "off" or _ocr_available is False:
        return ""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        with _ocr_lock:
            if _ocr_available is None:
                log.info("OCR of images is off: pytesseract is not installed")
            _ocr_available = False
        return ""
    try:
        with Image.open(io.BytesIO(data)) as im:
            text = pytesseract.image_to_string(im, timeout=30)
    except pytesseract.TesseractNotFoundError:
        with _ocr_lock:
            if _ocr_available is None:
                log.info("OCR of images is off: the tesseract program is not installed")
            _ocr_available = False
        return ""
    except Exception as exc:  # noqa: BLE001 - OCR is best effort
        log.warning("OCR of an image failed: %s", exc)
        return ""
    _ocr_available = True
    words = [w for w in text.split() if any(c.isalnum() for c in w)]
    return " ".join(words)[:1000] if len(words) >= 2 else ""


# ---------------------------------------------------------------------------- PDF

def _pdf_layout(page) -> tuple[dict[str, tuple[float, float, float, float]], list[tuple[float, str]]]:
    """Where each image XObject is drawn (x0, y0, x1, y1 in points) and the text lines with their height on the
    page, top to bottom."""
    placements: dict[str, tuple[float, float, float, float]] = {}
    fragments: list[tuple[float, float, str]] = []

    def before(op, args, cm, tm):
        if op == b"Do" and args:
            a, b, c, d, e, f = (float(v) for v in cm)
            xs, ys = (e, e + a, e + c, e + a + c), (f, f + b, f + d, f + b + d)
            placements.setdefault(str(args[0]).lstrip("/"), (min(xs), min(ys), max(xs), max(ys)))

    def on_text(text, cm, tm, font, size):
        if not text or not text.strip():
            return
        x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
        y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
        step = (float(size or 10) or 10) * 1.2
        for k, part in enumerate(p for p in text.split("\n") if p.strip()):
            fragments.append((float(y) - k * step, float(x), part))

    try:
        page.extract_text(visitor_operand_before=before, visitor_text=on_text)
    except Exception:  # noqa: BLE001 - no layout: images keep their page but get no nearby text
        log.debug("No layout for a PDF page", exc_info=True)
    lines: list[tuple[float, str]] = []
    for y, x, part in sorted(fragments, key=lambda f: (-f[0], f[1])):
        if lines and abs(lines[-1][0] - y) <= 2.0:
            lines[-1] = (lines[-1][0], f"{lines[-1][1]} {part}")
        else:
            lines.append((y, part))
    return placements, [(y, _clean(t)) for y, t in lines if _clean(t)]


def _pdf_candidates(path: Path, opts: ImageOptions, excel: dict) -> Iterator[_Candidate]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        reader.decrypt("")
    k = opts.context_lines
    for number, page in enumerate(reader.pages, start=1):
        try:
            images = list(page.images)
        except Exception as exc:  # noqa: BLE001 - one broken page must not stop the others
            log.warning("Could not list the images on page %s of %s: %s", number, path.name, exc)
            continue
        if not images:
            continue
        placements, lines = _pdf_layout(page)
        page_caption = [t for _, t in lines if CAPTION_RE.match(t)]
        for image in images:
            try:
                data = image.data
            except Exception as exc:  # noqa: BLE001
                log.warning("Could not read an image on page %s of %s: %s", number, path.name, exc)
                continue
            key = image.name.rsplit(".", 1)[0].split("/")[-1]
            box = placements.get(key)
            if box is None:  # inline image or drawn inside a form: page-level context only
                caption = page_caption[0] if len(page_caption) == 1 else ""
                yield _Candidate(data, f"Page {number}", number, image.name, caption=caption)
                continue
            x0, y0, x1, y1 = box
            above = [t for y, t in lines if y > y1][-k:] if k else []
            beside = [t for y, t in lines if y0 <= y <= y1][: 2 * k]
            below = [t for y, t in lines if y < y0][:k]
            # The closest caption: just below, then just above, then beside the image.
            caption = _caption_of(below) or _caption_of(list(reversed(above))) or _caption_of(beside)
            if caption in below:  # a caption that wraps: "Figure 1: ... barcode scanner" / "window"
                for line in below[below.index(caption) + 1:]:
                    if caption.endswith((".", ":", ";")) or not line[:1].islower():
                        break
                    caption = f"{caption} {line}"
            yield _Candidate(data, f"Page {number}", number, image.name, caption=caption,
                             nearby=above + beside + below, shown_size=(x1 - x0, y1 - y0))


# ---------------------------------------------------------------------------- Word

def _docx_candidates(path: Path, opts: ImageOptions, excel: dict) -> Iterator[_Candidate]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks: list[tuple[str, bool]] = []     # (text, is_caption) for every paragraph and table, in order
    refs: list[tuple[int, str, str, str, list[str]]] = []   # (block, section, rId, alt text, names)
    section = ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            text = _clean(paragraph.text)
            if text and (style == "Title" or style.startswith("Heading")):
                section = text
            blocks.append((text, style == "Caption" or bool(CAPTION_RE.match(text))))
        elif tag == "tbl":
            rows = [" | ".join(_clean(c.text) for c in row.cells) for row in Table(child, doc).rows]
            blocks.append((" ; ".join(r for r in rows if r.strip(" |")), False))
        else:
            continue
        index = len(blocks) - 1
        for drawing in child.iter(_q("w", "drawing")):
            doc_pr = next(drawing.iter(_q("wp", "docPr")), None)
            pic_pr = next(drawing.iter(_q("pic", "cNvPr")), None)
            blip = next(drawing.iter(_q("a", "blip")), None)
            rid = blip.get(_q("r", "embed")) if blip is not None else None
            if not rid:
                continue
            alt = _clean((doc_pr.get("descr") or doc_pr.get("title") or "") if doc_pr is not None else "")
            names = [e.get("name", "") for e in (pic_pr, doc_pr) if e is not None]
            refs.append((index, section, rid, alt, names))
        for shape in child.iter(_q("v", "imagedata")):   # legacy (VML) pictures
            rid = shape.get(_q("r", "id"))
            if rid:
                refs.append((index, section, rid, _clean(shape.get(_q("o", "title")) or ""), []))

    k = opts.context_lines
    for index, section, rid, alt, names in refs:
        part = doc.part.related_parts.get(rid)
        if part is None or not hasattr(part, "blob"):
            continue  # a linked (external) picture
        names.append(posixpath.basename(str(part.partname)))
        name = next((n for n in names if meaningful_name(n)), names[-1])
        before = [t for t, _ in blocks[max(0, index - k):index] if t]
        own = [blocks[index][0]] if blocks[index][0] else []
        after = [t for t, _ in blocks[index + 1:index + 1 + k] if t]
        near = [blocks[i] for i in (index + 1, index - 1, index, index + 2, index - 2) if 0 <= i < len(blocks)]
        caption = next((t for t, is_caption in near if is_caption and t), "")
        yield _Candidate(part.blob, f"Section: {section}" if section else "Document", None, name, alt, caption,
                         before + own + after)


# ---------------------------------------------------------------------------- Excel

def _rels(zf: zipfile.ZipFile, part: str) -> dict[str, str]:
    """Relationship id -> target part name for one part of an Office file (internal targets only)."""
    folder, base = posixpath.split(part)
    try:
        root = ET.fromstring(zf.read(posixpath.join(folder, "_rels", base + ".rels")))
    except KeyError:
        return {}
    out = {}
    for rel in root.iter(_q("rel", "Relationship")):
        target = rel.get("Target") or ""
        if rel.get("TargetMode") == "External" or not target:
            continue
        out[rel.get("Id")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(
            posixpath.join(folder, target))
    return out


def _xlsx_candidates(path: Path, opts: ImageOptions, excel: dict) -> Iterator[_Candidate]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from .extract import _cell_text

    skip_hidden = bool(excel.get("skip_hidden_sheets", True))
    k = opts.context_lines
    workbook = None
    with zipfile.ZipFile(path) as zf:
        sheets = _rels(zf, "xl/workbook.xml")
        root = ET.fromstring(zf.read("xl/workbook.xml"))
        try:
            for sheet in root.iter(_q("main", "sheet")):
                title = sheet.get("name") or ""
                if skip_hidden and (sheet.get("state") or "visible") != "visible":
                    continue
                sheet_part = sheets.get(sheet.get(_q("r", "id")) or "")
                if not sheet_part or sheet_part not in zf.namelist():
                    continue
                sheet_rels = _rels(zf, sheet_part)
                drawings = [sheet_rels.get(d.get(_q("r", "id")) or "")
                            for d in ET.fromstring(zf.read(sheet_part)).iter(_q("main", "drawing"))]
                for drawing_part in (d for d in drawings if d and d in zf.namelist()):
                    media = _rels(zf, drawing_part)
                    for anchor in ET.fromstring(zf.read(drawing_part)):
                        start = anchor.find(_q("xdr", "from"))
                        row = int(start.findtext(_q("xdr", "row")) or 0) + 1 if start is not None else 1
                        col = int(start.findtext(_q("xdr", "col")) or 0) + 1 if start is not None else 1
                        for pic in anchor.iter(_q("xdr", "pic")):
                            props = next(pic.iter(_q("xdr", "cNvPr")), None)
                            blip = next(pic.iter(_q("a", "blip")), None)
                            target = media.get(blip.get(_q("r", "embed")) or "") if blip is not None else None
                            if not target or target not in zf.namelist():
                                continue
                            if workbook is None:
                                workbook = load_workbook(str(path), read_only=True, data_only=True)
                            nearby = []
                            if title in workbook.sheetnames:
                                for values in workbook[title].iter_rows(min_row=max(1, row - k), max_row=row + k,
                                                                        values_only=True):
                                    cells = [_cell_text(v) for v in values if _cell_text(v)]
                                    if cells:
                                        nearby.append(" | ".join(cells))
                            names = [props.get("name", "") if props is not None else "", posixpath.basename(target)]
                            alt = _clean((props.get("descr") or props.get("title") or "") if props is not None else "")
                            yield _Candidate(zf.read(target), f"Sheet '{title}', cell {get_column_letter(col)}{row}",
                                             None, next((n for n in names if meaningful_name(n)), names[-1]), alt,
                                             _caption_of(nearby), [f"Sheet: {title}", *nearby])
        finally:
            if workbook is not None:
                workbook.close()


# ---------------------------------------------------------------------------- entry point

SOURCES = {"pdf": _pdf_candidates, "docx": _docx_candidates, "xlsx": _xlsx_candidates}


def extract_images(path: Path, file_type: str, opts: ImageOptions | None = None,
                   excel: dict | None = None) -> list[ExtractedImage]:
    """The document's images worth showing, in document order. Never raises: images are best effort."""
    opts = opts or ImageOptions()
    source = SOURCES.get(file_type)
    if not opts.enabled or source is None:
        return []
    if not pillow_available():
        log.warning("Image extraction is off: Pillow is not installed")
        return []
    kept: dict[str, ExtractedImage] = {}
    places: dict[str, set[str]] = {}
    try:
        for cand in source(Path(path), opts, excel or {}):
            if cand.shown_size and not opts.keeps(round(cand.shown_size[0]), round(cand.shown_size[1])):
                continue  # drawn as an icon, bullet or rule, whatever its pixel size
            decoded = _decode(cand.data)
            if decoded is None:
                continue
            data, fmt, width, height = decoded
            if not opts.keeps(width, height):
                continue
            image_id = hashlib.sha256(data).hexdigest()[:16]
            places.setdefault(image_id, set()).add(cand.location)
            if image_id in kept:
                kept[image_id].occurrences += 1
                continue
            if len(kept) >= opts.max_per_document:
                continue
            kept[image_id] = ExtractedImage(
                data, fmt, width, height, cand.location, cand.page, _clean(cand.original_name)[:200],
                _clean(cand.alt_text)[:500] if meaningful_name(cand.alt_text) else "", _clean(cand.caption)[:300],
                "\n".join(_clean(t)[:400] for t in cand.nearby if _clean(t)))
    except Exception as exc:  # noqa: BLE001 - a broken image part must not fail the document's text
        log.warning("Image extraction from %s stopped early: %s", Path(path).name, exc)
    images = [img for iid, img in kept.items() if not (opts.repeated_pages and len(places[iid]) >= opts.repeated_pages)]
    for image in images:
        image.ocr_text = ocr_text(image.data, opts.ocr)
    return images


# ---------------------------------------------------------------------------- store

class ImageStore:
    """Image files and their manifest in the images container: ``<doc_id>/<image_id>.<ext>`` and
    ``<doc_id>/images.json``."""

    def __init__(self, blobs, container: str):
        self.blobs = blobs
        self.container = container

    def save(self, doc_id: str, file_name: str, images: list[ExtractedImage]) -> list[dict]:
        """Replace the document's images; returns their metadata records."""
        self.delete_document(doc_id)
        records = []
        for image in images:
            name = f"{doc_id}/{image.image_id}.{image.extension}"
            self.blobs.put_bytes(self.container, name, image.data, image.content_type)
            records.append(image.record(doc_id, file_name, f"{self.container}/{name}"))
        if records:
            self.blobs.put_json(self.container, f"{doc_id}/{MANIFEST}",
                                {"docId": doc_id, "fileName": file_name, "images": records})
        return records

    def manifest(self, doc_id: str) -> list[dict]:
        try:
            data = self.blobs.get_json_or_none(self.container, f"{doc_id}/{MANIFEST}")
        except ValueError:  # not a valid blob name
            return []
        return list((data or {}).get("images") or [])

    def get(self, doc_id: str, image_id: str) -> dict | None:
        if not IMAGE_ID_RE.fullmatch(image_id or ""):
            return None
        return next((r for r in self.manifest(doc_id) if r.get("imageId") == image_id), None)

    def iter_bytes(self, record: dict):
        ext = "jpg" if record.get("format") == "jpeg" else record.get("format")
        return self.blobs.iter_bytes(self.container, f"{record['docId']}/{record['imageId']}.{ext}")

    def delete_document(self, doc_id: str) -> int:
        blobs = self.blobs.list(self.container, prefix=f"{doc_id}/")
        for blob in blobs:
            self.blobs.delete(self.container, blob.name)
        return sum(1 for b in blobs if not b.name.endswith(MANIFEST))

    def clear(self) -> None:
        for blob in self.blobs.list(self.container):
            self.blobs.delete(self.container, blob.name)
