"""Text extraction for PDF (pypdf), Word (python-docx) and Excel (openpyxl), yielding passages lazily.

Each extractor works page by page, section by section or sheet by sheet, so only one part of a file is
turned into text at a time. There is no OCR of pages: an image-only PDF yields no passages. Embedded images are
extracted separately (images.py).
"""

from __future__ import annotations

import datetime as dt
import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from .chunk import ChunkOptions, chunk_rows, chunk_text

log = logging.getLogger(__name__)

FILE_TYPES = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx"}
FILE_TYPE_LABELS = {"pdf": "PDF", "docx": "Word", "xlsx": "Excel"}
CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class ExtractionError(Exception):
    """The file could not be read (corrupt, encrypted or not what its extension says)."""


@dataclass(frozen=True)
class Passage:
    text: str
    location: str
    image_id: str = ""   # set for the passage that describes an extracted image (images.py)


def file_type_for(name: str) -> str | None:
    return FILE_TYPES.get(Path(name).suffix.lower())


def extract_passages(path: Path, file_type: str, opts: ChunkOptions | None = None,
                     excel: dict | None = None) -> Iterator[Passage]:
    opts = opts or ChunkOptions()
    if file_type == "pdf":
        return _pdf(path, opts)
    if file_type == "docx":
        return _docx(path, opts)
    if file_type == "xlsx":
        return _xlsx(path, opts, excel or {})
    raise ExtractionError(f"Unsupported file type: {file_type}")


# ---------------------------------------------------------------------------- PDF

def _pdf(path: Path, opts: ChunkOptions) -> Iterator[Passage]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                raise ExtractionError("This PDF is password protected.") from None
        pages = len(reader.pages)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, OSError, KeyError) as exc:
        raise ExtractionError(f"This PDF could not be read ({exc.__class__.__name__}).") from exc

    for number in range(pages):
        try:
            text = reader.pages[number].extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one broken page must not stop the others
            log.warning("Could not read page %s of %s: %s", number + 1, path.name, exc)
            continue
        for passage in chunk_text(text, opts):
            yield Passage(passage, f"Page {number + 1}")


# ---------------------------------------------------------------------------- Word

def _docx(path: Path, opts: ChunkOptions) -> Iterator[Passage]:
    import zipfile

    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(str(path))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError("This Word file could not be read.") from exc

    headings: list[tuple[int, str]] = []
    buffer: list[str] = []

    def heading_level(paragraph) -> int | None:
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style == "Title":
            return 0
        if style.startswith("Heading"):
            try:
                return int(style.split()[-1])
            except ValueError:
                return 1
        return None

    def flush():
        text = "\n\n".join(buffer).strip()
        buffer.clear()
        if not text:
            return
        path_text = " > ".join(h for _, h in headings)
        location = f"Section: {headings[-1][1]}" if headings else "Document"
        for passage in chunk_text(text, opts, prefix=path_text):
            yield Passage(passage, location)

    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, doc)
            text = paragraph.text.strip()
            if not text:
                continue
            level = heading_level(paragraph)
            if level is not None:
                yield from flush()
                headings[:] = [h for h in headings if h[0] < level] + [(level, text)]
            else:
                buffer.append(text)
        elif tag == "tbl":
            rows = []
            for row in Table(child, doc).rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or value != cells[-1]:  # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                buffer.append("\n".join(rows))
    yield from flush()


# ---------------------------------------------------------------------------- Excel

def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        return f"{int(value)}" if value.is_integer() else f"{value:.4g}" if abs(value) < 1e-3 else f"{round(value, 4)}"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    return " ".join(str(value).split())


def _xlsx(path: Path, opts: ChunkOptions, excel: dict) -> Iterator[Passage]:
    import zipfile

    from openpyxl import load_workbook

    skip_hidden = bool(excel.get("skip_hidden_sheets", True))
    max_rows = int(excel.get("max_rows_per_sheet", 20000))
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractionError("This Excel file could not be read.") from exc
    try:
        for sheet in workbook.worksheets:
            if skip_hidden and getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            header: list[str] | None = None
            rows: list[tuple[int, str]] = []
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if number > max_rows:
                    log.info("Sheet %s of %s truncated at %s rows", sheet.title, path.name, max_rows)
                    break
                cells = [_cell_text(v) for v in values]
                while cells and not cells[-1]:
                    cells.pop()
                if not any(cells):
                    continue  # skip empty rows
                if header is None:
                    header = cells
                    continue
                rows.append((number, " | ".join(cells)))
            if header is None:
                continue
            if not rows:  # a sheet with only a header still carries meaning
                yield Passage(f"Sheet: {sheet.title}\n" + " | ".join(header), f"Sheet '{sheet.title}', row 1")
                continue
            prefix = f"Sheet: {sheet.title}\nColumns: " + " | ".join(header)
            for first, last, text in chunk_rows(rows, opts, prefix=prefix):
                where = f"row {first}" if first == last else f"rows {first}-{last}"
                yield Passage(text, f"Sheet '{sheet.title}', {where}")
    finally:
        workbook.close()
